from io import BytesIO
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter


# ===== Estilos reutilizables =====
HEADER_FILL = PatternFill('solid', fgColor='111827')
HEADER_FONT = Font(color='FFFFFF', bold=True, size=11)
SUBHEADER_FILL = PatternFill('solid', fgColor='F3F4F6')
SUBHEADER_FONT = Font(bold=True, size=10)
TITLE_FONT = Font(bold=True, size=14, color='111827')
BORDER = Border(
    left=Side(style='thin', color='E5E7EB'),
    right=Side(style='thin', color='E5E7EB'),
    top=Side(style='thin', color='E5E7EB'),
    bottom=Side(style='thin', color='E5E7EB'),
)


def _ajustar_columnas(ws, min_width=10, max_width=40):
    """Auto-ajusta el ancho de las columnas según el contenido."""
    for col in ws.columns:
        length = 0
        letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                length = max(length, len(str(cell.value)))
        adjusted = min(max(length + 2, min_width), max_width)
        ws.column_dimensions[letter].width = adjusted


def _escribir_titulo(ws, titulo, subtitulo='', fila=1):
    """Escribe el título y subtítulo de la hoja."""
    ws.cell(row=fila, column=1, value=titulo).font = TITLE_FONT
    ws.merge_cells(
        start_row=fila, start_column=1, end_row=fila, end_column=6
    )
    if subtitulo:
        ws.cell(row=fila + 1, column=1, value=subtitulo).font = Font(
            italic=True, color='6B7280', size=10
        )
        ws.merge_cells(
            start_row=fila + 1, start_column=1, end_row=fila + 1, end_column=6
        )
        return fila + 3
    return fila + 2


def _escribir_headers(ws, headers, fila):
    """Escribe los headers de la tabla con estilos."""
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=fila, column=col_num, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = BORDER
    ws.row_dimensions[fila].height = 22
    return fila + 1


def _escribir_fila(ws, valores, fila, alinear_derecha_cols=None):
    """Escribe una fila de datos con estilos."""
    alinear_derecha_cols = alinear_derecha_cols or []
    for col_num, valor in enumerate(valores, 1):
        cell = ws.cell(row=fila, column=col_num, value=valor)
        cell.border = BORDER
        if col_num in alinear_derecha_cols:
            cell.alignment = Alignment(horizontal='right')
    return fila + 1


# ============================================================
# REPORTE: ALUMNOS
# ============================================================
def exportar_alumnos_excel(alumnos):
    """Exporta listado de alumnos a Excel con formato."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Alumnos'

    fila = _escribir_titulo(
        ws, 'Listado de Alumnos',
        f'Generado el {datetime.now().strftime("%d/%m/%Y %H:%M")} · '
        f'{alumnos.count()} alumnos'
    )

    headers = [
        'Apellido', 'Nombre', 'RUT', 'Fecha nac.', 'Edad',
        'Cinturón', 'Grados', 'Categoría', 'Estado',
        'Teléfono', 'Email', 'Tutor', 'Tel. tutor', 'Fecha inicio',
    ]
    fila = _escribir_headers(ws, headers, fila)

    for alumno in alumnos:
        edad = alumno.calcular_edad() if hasattr(alumno, 'calcular_edad') else ''
        fila = _escribir_fila(ws, [
            alumno.apellido,
            alumno.nombre,
            alumno.rut,
            alumno.fecha_nacimiento.strftime('%d/%m/%Y'),
            edad,
            alumno.get_cinturon_actual_display(),
            f'{alumno.grados_cinturon}º',
            alumno.get_categoria_display(),
            alumno.get_estado_display(),
            alumno.telefono or '—',
            alumno.email or '—',
            alumno.tutor_nombre or '—',
            alumno.tutor_telefono or '—',
            alumno.fecha_inicio.strftime('%d/%m/%Y'),
        ])

    _ajustar_columnas(ws)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ============================================================
# REPORTE: RECAUDACIÓN
# ============================================================
def exportar_recaudacion_excel(pagos, titulo, subtitulo=''):
    """Exporta reporte de recaudación a Excel con hoja de desglose."""
    wb = Workbook()

    # ===== Hoja 1: Detalle =====
    ws = wb.active
    ws.title = 'Detalle'
    fila = _escribir_titulo(ws, titulo, subtitulo)

    headers = [
        'Fecha', 'N° Recibo', 'Alumno', 'Período', 'Plan',
        'Método de pago', 'Monto base', 'Descuento', 'Total',
    ]
    fila = _escribir_headers(ws, headers, fila)

    total_general = Decimal('0')
    for p in pagos:
        fila = _escribir_fila(ws, [
            p.fecha_pago.strftime('%d/%m/%Y') if p.fecha_pago else '—',
            p.numero_recibo or '—',
            f'{p.alumno.apellido}, {p.alumno.nombre}',
            p.periodo_display,
            p.plan.nombre if p.plan else '—',
            p.get_metodo_pago_display() or '—',
            float(p.monto),
            float(p.descuento),
            float(p.monto_final),
        ], alinear_derecha_cols=[7, 8, 9])
        total_general += p.monto_final

    # Fila de total
    fila += 1
    ws.cell(row=fila, column=8, value='TOTAL').font = Font(bold=True, size=12)
    ws.cell(row=fila, column=9, value=float(total_general)).font = Font(
        bold=True, size=12, color='10B981'
    )
    ws.cell(row=fila, column=9).alignment = Alignment(horizontal='right')

    _ajustar_columnas(ws)

    # ===== Hoja 2: Desglose por método =====
    ws2 = wb.create_sheet('Desglose')
    fila2 = _escribir_titulo(ws2, 'Desglose por método de pago')

    headers2 = ['Método', 'Cantidad', 'Total']
    fila2 = _escribir_headers(ws2, headers2, fila2)

    from django.db.models import Sum, Count
    por_metodo = pagos.values('metodo_pago').annotate(
        total=Sum('monto_final'), cantidad=Count('id')
    ).order_by('-total')

    metodos_map = dict(Pago_Metodo_choices())
    for item in por_metodo:
        fila2 = _escribir_fila(ws2, [
            metodos_map.get(item['metodo_pago'], item['metodo_pago'] or '—'),
            item['cantidad'],
            float(item['total']),
        ], alinear_derecha_cols=[3])

    _ajustar_columnas(ws2)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def Pago_Metodo_choices():
    from pagos.models import Pago
    return Pago.Metodo.choices


# ============================================================
# REPORTE: ASISTENCIA
# ============================================================
def exportar_asistencia_excel(asistencias, titulo, subtitulo=''):
    """Exporta registros de asistencia a Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Asistencia'

    fila = _escribir_titulo(ws, titulo, subtitulo)

    headers = ['Fecha', 'Alumno', 'RUT', 'Clase', 'Hora', 'Estado']
    fila = _escribir_headers(ws, headers, fila)

    for a in asistencias:
        fila = _escribir_fila(ws, [
            a.fecha.strftime('%d/%m/%Y'),
            f'{a.alumno.apellido}, {a.alumno.nombre}',
            a.alumno.rut,
            a.clase.nombre,
            a.hora_registro.strftime('%H:%M'),
            'Presente' if a.presente else 'Ausente',
        ])

    _ajustar_columnas(ws)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ============================================================
# REPORTE: APTOS PARA EXAMEN
# ============================================================
def exportar_aptos_examen_excel(alumnos, fila_datos):
    """
    Exporta listado de alumnos aptos para examen de cinturón.
    `alumnos` es la lista ya calculada con estadísticas.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = 'Aptos para examen'

    fila = _escribir_titulo(
        ws, 'Alumnos aptos para examen de cinturón',
        f'Generado el {datetime.now().strftime("%d/%m/%Y %H:%M")}'
    )

    headers = [
        'Alumno', 'RUT', 'Cinturón', 'Grados',
        'Asistencias (90d)', 'Total asistencias',
        'Certificado médico', 'Apto', 'Observaciones',
    ]
    fila = _escribir_headers(ws, headers, fila)

    for item in alumnos:
        alumno = item['alumno']
        fila = _escribir_fila(ws, [
            f'{alumno.apellido}, {alumno.nombre}',
            alumno.rut,
            alumno.get_cinturon_actual_display(),
            f'{alumno.grados_cinturon}º',
            item['asistencias_90d'],
            item['total_asistencias'],
            item['certificado_estado'],
            'Sí' if item['apto'] else 'No',
            item['observaciones'],
        ])

    _ajustar_columnas(ws)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer